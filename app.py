import time
import io
import streamlit as st
from dotenv import load_dotenv
from docx import Document
from fpdf import FPDF

from utils.audio_processor import process_input
from core.transcriber import transcribe_all
from core.summarizer import summarize, generate_title
from core.extractor import extract_action_items, extract_key_decisions, extract_questions
from core.rag_engine import build_rag_chain, ask_question

load_dotenv()

st.set_page_config(page_title="AI Study Agent", page_icon="📚", layout="wide")

# ---------------------------------------------------------------------------
# Design tokens & global styling
# ---------------------------------------------------------------------------
st.markdown(
    """
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link href="https://fonts.googleapis.com/css2?family=Fraunces:opsz,wght@9..144,400;9..144,500;9..144,600&family=IBM+Plex+Sans:wght@400;500;600&display=swap" rel="stylesheet">

    <style>
    :root {
        --ink: #1C1B19;
        --ink-soft: #26241F;
        --paper: #EAE1CE;
        --paper-soft: #F3EEE3;
        --brass: #D6B23E;
        --teal: #2F5D5A;
        --line: #4B453B;
        --text-on-ink: #F3EEE3;
        --text-on-paper: #241F18;
    }

    html, body, [class*="css"] { font-family: 'IBM Plex Sans', sans-serif; }

    .stApp { background-color: var(--ink); color: var(--text-on-ink); }

    /* ---- Sidebar: the "slate" ---- */
    section[data-testid="stSidebar"] {
        background-color: var(--ink-soft);
        border-right: 1px solid var(--line);
    }
    section[data-testid="stSidebar"] * { color: var(--text-on-ink) !important; }
    section[data-testid="stSidebar"] h1 {
        font-family: 'Fraunces', serif;
        font-weight: 500;
        font-size: 1.7rem;
        letter-spacing: 0.01em;
        border-bottom: 1px solid var(--line);
        padding-bottom: 0.6rem;
        margin-bottom: 0.2rem;
    }
    section[data-testid="stSidebar"] input, section[data-testid="stSidebar"] textarea {
        background-color: var(--paper) !important;
        color: var(--text-on-paper) !important;
        border-radius: 2px !important;
        border: 1px solid var(--line) !important;
    }
    section[data-testid="stSidebar"] [data-baseweb="select"] > div {
        background-color: var(--paper) !important;
        color: var(--text-on-paper) !important;
        border-radius: 2px !important;
    }

    /* ---- Buttons: flat, brass, no rounded-pill default ---- */
    .stButton > button {
        background-color: var(--brass);
        color: var(--ink);
        border: none;
        border-radius: 2px;
        font-weight: 600;
        padding: 0.55rem 1.1rem;
        transition: background-color 0.15s ease;
    }
    .stButton > button:hover { background-color: #DCA45A; color: var(--ink); }
    .stButton > button:disabled { background-color: #6B6152; color: #B9B0A0; }

    /* ---- Download button: secondary style (teal, outlined) ---- */
    .stDownloadButton > button {
        background-color: transparent;
        color: var(--text-on-ink);
        border: 1px solid var(--teal);
        border-radius: 2px;
        font-weight: 600;
    }
    .stDownloadButton > button:hover {
        background-color: var(--teal);
        color: var(--text-on-ink);
    }

    /* ---- Title card hero ---- */
    .title-card {
        border-top: 2px solid var(--brass);
        border-bottom: 1px solid var(--line);
        padding: 1.4rem 0 1.2rem 0;
        margin-bottom: 1.8rem;
    }
    .title-card .kicker {
        font-size: 0.85rem;
        color: #B9AE94;
        margin-bottom: 0.3rem;
    }
    .title-card h1 {
        font-family: 'Fraunces', serif;
        font-weight: 500;
        font-size: 2.4rem;
        line-height: 1.15;
        margin: 0;
        color: var(--text-on-ink);
    }

    /* ---- Step checklist (replaces generic progress bar) ---- */
    .step-row {
        display: flex;
        align-items: baseline;
        gap: 0.7rem;
        padding: 0.35rem 0;
        border-bottom: 1px dashed var(--line);
        font-size: 0.95rem;
    }
    .step-row .mark { width: 1.2rem; color: var(--brass); }
    .step-row.pending { color: #8A8171; }
    .step-row.done .mark { color: #7FA88A; }
    .step-row.active .mark { color: var(--brass); }

    /* ---- Tabs restyled as a dossier index, not pills ---- */
    .stTabs [data-baseweb="tab-list"] {
        gap: 1.6rem;
        border-bottom: 1px solid var(--line);
    }
    .stTabs [data-baseweb="tab"] {
        background-color: transparent;
        color: #B9AE94;
        font-weight: 500;
        padding-bottom: 0.6rem;
    }
    .stTabs [aria-selected="true"] {
        color: var(--text-on-ink) !important;
        border-bottom: 2px solid var(--brass) !important;
    }

    /* ---- Dossier body text ---- */
    .dossier-body {
        background-color: var(--paper);
        color: var(--text-on-paper);
        padding: 1.4rem 1.6rem;
        border-radius: 2px;
        line-height: 1.6;
    }

    /* ---- Chat ---- */
    [data-testid="stChatMessage"] {
        background-color: transparent;
        border-bottom: 1px solid var(--line);
        border-radius: 0;
        padding-bottom: 0.8rem;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

STEPS = [
    "Pulling in the lecture",
    "Transcribing the recording",
    "Drafting a title",
    "Writing the summary",
    "Listing action items",
    "Logging key decisions",
    "Flagging open questions",
    "Setting up the study chat",
]

def build_docx_bytes(result: dict) -> bytes:
    """Build a Word document with the notes and return it as bytes."""
    doc = Document()
    doc.add_heading(result["title"], level=1)

    sections = [
        ("Summary", result["summary"]),
        ("Action Items", result["action_items"]),
        ("Key Decisions", result["key_decisions"]),
        ("Open Questions", result["open_questions"]),
        ("Full Transcript", result["transcript"]),
    ]
    for heading, body in sections:
        doc.add_heading(heading, level=2)
        for line in body.split("\n"):
            doc.add_paragraph(line)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_pdf_bytes(result: dict) -> bytes:
    """Build a simple PDF with the notes and return it as bytes."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, result["title"].encode("latin-1", "replace").decode("latin-1"))
    pdf.ln(2)

    sections = [
        ("Summary", result["summary"]),
        ("Action Items", result["action_items"]),
        ("Key Decisions", result["key_decisions"]),
        ("Open Questions", result["open_questions"]),
        ("Full Transcript", result["transcript"]),
    ]
    for heading, body in sections:
        pdf.set_font("Helvetica", "B", 13)
        pdf.multi_cell(0, 8, heading)
        pdf.ln(1)
        pdf.set_font("Helvetica", "", 11)
        safe_body = body.encode("latin-1", "replace").decode("latin-1")
        pdf.multi_cell(0, 6, safe_body)
        pdf.ln(4)

    return bytes(pdf.output(dest="S"))


# ---------------------------------------------------------------------------
# Session state
# ---------------------------------------------------------------------------
if "result" not in st.session_state:
    st.session_state.result = None
if "chat_history" not in st.session_state:
    st.session_state.chat_history = []
if "processing" not in st.session_state:
    st.session_state.processing = False


def render_steps(placeholder, current_index, total=len(STEPS)):
    """Render a checklist: done steps get a check, current gets a dot, rest are dimmed."""
    rows = []
    for i, label in enumerate(STEPS):
        if i < current_index:
            rows.append(f'<div class="step-row done"><span class="mark">✓</span>{label}</div>')
        elif i == current_index:
            rows.append(f'<div class="step-row active"><span class="mark">●</span>{label}…</div>')
        else:
            rows.append(f'<div class="step-row pending"><span class="mark">·</span>{label}</div>')
    placeholder.markdown("".join(rows), unsafe_allow_html=True)


def run_pipeline_with_progress(source: str, language: str, placeholder) -> dict:
    render_steps(placeholder, 0)
    chunks = process_input(source)

    render_steps(placeholder, 1)
    transcript = transcribe_all(chunks, language)

    render_steps(placeholder, 2)
    title = generate_title(transcript)
    time.sleep(5)

    render_steps(placeholder, 3)
    summary = summarize(transcript)
    time.sleep(5)

    render_steps(placeholder, 4)
    action_items = extract_action_items(transcript)
    time.sleep(5)

    render_steps(placeholder, 5)
    decisions = extract_key_decisions(transcript)
    time.sleep(5)

    render_steps(placeholder, 6)
    questions = extract_questions(transcript)
    time.sleep(5)

    render_steps(placeholder, 7)
    rag_chain = build_rag_chain(transcript)

    render_steps(placeholder, 8)

    return {
        "title": title,
        "transcript": transcript,
        "summary": summary,
        "action_items": action_items,
        "key_decisions": decisions,
        "open_questions": questions,
        "rag_chain": rag_chain,
    }


# ---------------------------------------------------------------------------
# Sidebar — the slate
# ---------------------------------------------------------------------------
with st.sidebar:
    st.title("📚 AI Study Agent")
    st.caption("Turn a lecture or video into a title, summary, key points, and a chat you can quiz yourself with.")

    source = st.text_input(
        "Lecture link or file path",
        placeholder="https://www.youtube.com/watch?v=...",
    )
    language = st.radio("Spoken language", ["english", "hinglish"], horizontal=True)

    run_clicked = st.button(
        "Build my notes",
        use_container_width=True,
        disabled=st.session_state.processing or not source.strip(),
    )

    if st.session_state.result is not None:
        st.divider()
        if st.button("Start a new lecture", use_container_width=True):
            st.session_state.result = None
            st.session_state.chat_history = []
            st.rerun()

# ---------------------------------------------------------------------------
# Run pipeline
# ---------------------------------------------------------------------------
if run_clicked and source.strip():
    st.session_state.processing = True
    st.session_state.chat_history = []

    st.markdown('<div class="title-card"><div class="kicker">Now studying</div><h1>Building your notes…</h1></div>', unsafe_allow_html=True)
    steps_placeholder = st.empty()

    try:
        st.session_state.result = run_pipeline_with_progress(source.strip(), language, steps_placeholder)
        st.toast("Notes are ready.")
    except Exception as e:
        st.error(f"Something interrupted the lecture: {e}")
        st.session_state.result = None
    finally:
        st.session_state.processing = False
        steps_placeholder.empty()
        st.rerun()

# ---------------------------------------------------------------------------
# Main content
# ---------------------------------------------------------------------------
if st.session_state.result is None and not st.session_state.processing:
    st.markdown(
        '<div class="title-card"><div class="kicker">Study desk — empty</div>'
        '<h1>Bring in a lecture to start your notes.</h1></div>',
        unsafe_allow_html=True,
    )
    st.write("Paste a lecture link or a local file path on the left, then press **Build my notes**.")

elif st.session_state.result is not None:
    result = st.session_state.result

    st.markdown(
        f'<div class="title-card"><div class="kicker">Today\'s notes</div><h1>{result["title"]}</h1></div>',
        unsafe_allow_html=True,
    )

    notes_text = (
        f"{result['title']}\n"
        f"{'=' * len(result['title'])}\n\n"
        f"SUMMARY\n-------\n{result['summary']}\n\n"
        f"ACTION ITEMS\n------------\n{result['action_items']}\n\n"
        f"KEY DECISIONS\n-------------\n{result['key_decisions']}\n\n"
        f"OPEN QUESTIONS\n--------------\n{result['open_questions']}\n\n"
        f"FULL TRANSCRIPT\n---------------\n{result['transcript']}\n"
    )
    safe_name = result["title"][:60].strip() or "study_notes"

    dl_col1, dl_col2, dl_col3 = st.columns(3)
    with dl_col1:
        st.download_button(
            "Download as .txt",
            data=notes_text,
            file_name=f"{safe_name}.txt",
            mime="text/plain",
            use_container_width=True,
        )
    with dl_col2:
        st.download_button(
            "Download as Word (.docx)",
            data=build_docx_bytes(result),
            file_name=f"{safe_name}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )
    with dl_col3:
        st.download_button(
            "Download as PDF",
            data=build_pdf_bytes(result),
            file_name=f"{safe_name}.pdf",
            mime="application/pdf",
            use_container_width=True,
        )

    tab_summary, tab_actions, tab_decisions, tab_questions, tab_transcript, tab_chat = st.tabs(
        ["Summary", "Action items", "Key decisions", "Open questions", "Full transcript", "Study chat"]
    )

    with tab_summary:
        st.markdown(f'<div class="dossier-body">{result["summary"]}</div>', unsafe_allow_html=True)

    with tab_actions:
        st.markdown(f'<div class="dossier-body">{result["action_items"]}</div>', unsafe_allow_html=True)

    with tab_decisions:
        st.markdown(f'<div class="dossier-body">{result["key_decisions"]}</div>', unsafe_allow_html=True)

    with tab_questions:
        st.markdown(f'<div class="dossier-body">{result["open_questions"]}</div>', unsafe_allow_html=True)

    with tab_transcript:
        st.text_area("Transcript", result["transcript"], height=420, label_visibility="collapsed")

    with tab_chat:
        st.caption("Ask about this lecture — answers stay grounded in the transcript.")

        for role, msg in st.session_state.chat_history:
            with st.chat_message(role):
                st.markdown(msg)

        question = st.chat_input("Ask something about this lecture...")
        if question:
            st.session_state.chat_history.append(("user", question))
            with st.chat_message("user"):
                st.markdown(question)

            with st.chat_message("agent"):
                with st.spinner("Checking the transcript..."):
                    answer = ask_question(result["rag_chain"], question)
                st.markdown(answer)
            st.session_state.chat_history.append(("agent", answer))